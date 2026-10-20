"""Slide generation (Groq) and PPTX building, with optional images."""

from __future__ import annotations

import io
import json
import os
from dataclasses import dataclass, field

from groq import Groq
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ============================================================================
# Data Models
# ============================================================================

@dataclass
class Slide:
    title: str
    bullets: list[str] = field(default_factory=list)
    notes: str = ""
    image_id: str = ""            # id chosen by the model (e.g. "p3_1", "w2", "s54")
    image: bytes | None = None    # real image bytes, filled in by the server


@dataclass
class Deck:
    deck_title: str
    subtitle: str
    slides: list[Slide] = field(default_factory=list)
    source: str = ""              # shown small on the title slide


# ============================================================================
# Slide Generation (Groq)
# ============================================================================

DEFAULT_MODEL = os.environ.get("GROQ_MODEL", "llama-3.3-70b-versatile")

GROQ_API_KEYS = [
    os.environ.get("GROQ_API_KEY"),
    os.environ.get("GROQ_API_KEY_2"),
    os.environ.get("GROQ_API_KEY_3"),
]

SYSTEM_PROMPT = """You are an expert presentation writer.
You will receive SOURCE MATERIAL (from a PDF or web page) and a TOPIC.
Produce a presentation about the TOPIC grounded in the SOURCE MATERIAL.

Rules:
- Only use facts found in the source material. Do not invent data, numbers, names, or whole sections.
- Be specific: keep drug names, doses, numbers, and study authors/years from the source. No filler bullets.
- Concise slide titles and 3-4 bullets per slide. Each bullet under 15 words.
- The user gives an approximate number of content slides. Treat it as a MAXIMUM. Always produce at least 4 content slides. The "slides" array must never be empty.
- Follow the logical order of the source.
- Image markers look like [IMAGE w3: description]. Set "image" to that id if it belongs on the slide, else empty string. One image per slide max, each id once only.
- Keep notes to 1 short sentence.

Return ONLY valid minified JSON:
{"deck_title":"string","subtitle":"string","slides":[{"title":"string","bullets":["string"],"notes":"string","image":"string"}]}"""


def _call_groq(user_prompt: str, model: str, api_key: str) -> str:
    resp = Groq(api_key=api_key).chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": user_prompt},
        ],
        response_format={"type": "json_object"},
        temperature=0.2,
        max_tokens=4096,
    )
    raw = (resp.choices[0].message.content or "").strip()
    if not raw:
        raise RuntimeError("Groq returned an empty response.")
    return raw


def generate_deck_groq(
    source_text: str,
    topic: str,
    num_slides: int,
    model: str = DEFAULT_MODEL,
    valid_image_ids: set[str] | None = None,
    _retry: bool = True,
) -> Deck:
    """Generate a deck using Groq's hosted LLMs with 3-key rotation."""
    user_prompt = (
        f"TOPIC: {topic}\n"
        f"MAXIMUM NUMBER OF CONTENT SLIDES: {num_slides}\n\n"
        f"SOURCE MATERIAL:\n{source_text}"
    )

    valid_keys = [k for k in GROQ_API_KEYS if k]
    if not valid_keys:
        raise RuntimeError("No GROQ_API_KEY found in environment.")

    last_exc = None
    for attempt, api_key in enumerate(valid_keys):
        try:
            raw = _call_groq(user_prompt, model, api_key)
            data = json.loads(raw)

            valid = valid_image_ids or set()
            used: set[str] = set()
            slides: list[Slide] = []
            for i, s in enumerate(data.get("slides", [])):
                if not isinstance(s, dict):
                    continue
                img = str(s.get("image", "") or "").strip()
                if img not in valid or img in used:
                    img = ""
                if img:
                    used.add(img)
                slides.append(
                    Slide(
                        title=str(s.get("title", "")).strip() or f"Slide {i + 1}",
                        bullets=[str(b).strip() for b in s.get("bullets", []) if str(b).strip()],
                        notes=str(s.get("notes", "")).strip(),
                        image_id=img,
                    )
                )

            if not slides:
                if _retry and attempt == len(valid_keys) - 1:
                    nudge = (
                        "IMPORTANT: your previous answer had an empty slides array. "
                        "Return at least 4 content slides built from the source below.\n\n"
                    )
                    return generate_deck_groq(
                        nudge + source_text, topic, num_slides, model, valid_image_ids, _retry=False
                    )
                continue

            return Deck(
                deck_title=str(data.get("deck_title", topic)).strip() or topic,
                subtitle=str(data.get("subtitle", "")).strip(),
                slides=slides,
            )

        except json.JSONDecodeError as exc:
            last_exc = RuntimeError(f"Model returned invalid JSON: {exc}")
        except Exception as exc:
            last_exc = RuntimeError(f"Groq call failed: {exc}")

    raise RuntimeError(f"All {len(valid_keys)} API keys failed. Last error: {last_exc}")


# ============================================================================
# PPTX Building
# ============================================================================

INK = RGBColor(0x1F, 0x29, 0x37)      # dark slate for text
ACCENT = RGBColor(0x25, 0x63, 0xEB)    # blue accent
MUTED = RGBColor(0x6B, 0x72, 0x80)     # gray for subtitles/footers
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _add_footer(slide, page_no: int):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(12.3), Inches(0.35))
    tf = box.text_frame
    tf.text = str(page_no)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.runs[0]
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED


def _add_accent_bar(slide, top: Emu, left: Emu = Inches(0.5), width: Emu = Inches(1.2)):
    bar = slide.shapes.add_shape(1, left, top, width, Inches(0.08))  # 1 = rectangle
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def _add_picture_fit(slide, data: bytes, left, top, box_w, box_h) -> None:
    """Place an image inside a box, keeping its aspect ratio and centering it."""
    with Image.open(io.BytesIO(data)) as im:
        w, h = im.size
    scale = min(box_w / w, box_h / h)
    pw, ph = int(w * scale), int(h * scale)
    slide.shapes.add_picture(
        io.BytesIO(data),
        int(left + (box_w - pw) / 2),
        int(top + (box_h - ph) / 2),
        width=pw,
        height=ph,
    )


def _title_slide(prs: Presentation, deck: Deck) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = INK

    _add_accent_bar(slide, top=Inches(2.6), left=Inches(1.0), width=Inches(1.6))

    title = slide.shapes.add_textbox(Inches(1.0), Inches(2.85), Inches(11.3), Inches(1.8))
    tf = title.text_frame
    tf.word_wrap = True
    tf.text = deck.deck_title
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = "Calibri"

    if deck.subtitle:
        sub = slide.shapes.add_textbox(Inches(1.0), Inches(4.7), Inches(11.3), Inches(0.9))
        tf = sub.text_frame
        tf.word_wrap = True
        tf.text = deck.subtitle
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(0xD1, 0xD5, 0xDB)
        run.font.name = "Calibri"

    if deck.source:
        src = slide.shapes.add_textbox(Inches(1.0), Inches(6.7), Inches(11.3), Inches(0.4))
        tf = src.text_frame
        tf.word_wrap = True
        tf.text = f"Source: {deck.source}"
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(11)
        run.font.color.rgb = MUTED
        run.font.name = "Calibri"


def _content_slide(prs: Presentation, slide_data, page_no: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    has_image = bool(slide_data.image)

    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.45), Inches(12.3), Inches(1.0))
    tf = title.text_frame
    tf.word_wrap = True
    tf.text = slide_data.title
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = INK
    run.font.name = "Calibri"

    _add_accent_bar(slide, top=Inches(1.45))

    body_w = Inches(6.4) if has_image else Inches(12.0)
    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.85), body_w, Inches(4.9))
    tf = body.text_frame
    tf.word_wrap = True

    for i, bullet in enumerate(slide_data.bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(12)
        p.font.size = Pt(18)
        p.font.color.rgb = INK
        p.font.name = "Calibri"

    if has_image:
        try:
            _add_picture_fit(
                slide, slide_data.image, Inches(7.3), Inches(1.85), Inches(5.5), Inches(4.9)
            )
        except Exception:
            pass  # a bad image should never break the whole deck

    if slide_data.notes:
        slide.notes_slide.notes_text_frame.text = slide_data.notes

    _add_footer(slide, page_no)


def build_pptx(deck: Deck) -> bytes:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    _title_slide(prs, deck)
    for i, slide_data in enumerate(deck.slides, start=2):
        _content_slide(prs, slide_data, i)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()